"""
Export Service - PDF and Excel report generation
Uses ReportLab for PDF and openpyxl for Excel
"""

from typing import Optional, Dict, Any, List
from datetime import datetime
import io


class ExportService:
    """Service for generating exportable reports"""

    async def generate_pdf(
        self, report_type: str, data: Dict[str, Any], title: Optional[str] = None
    ) -> bytes:
        """Generate a PDF report"""
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import mm
            from reportlab.platypus import (
                SimpleDocTemplate,
                Paragraph,
                Spacer,
                Table,
                TableStyle,
                Image,
            )
            from reportlab.lib.enums import TA_CENTER, TA_LEFT
        except ImportError:
            raise ImportError(
                "reportlab is required for PDF generation. Install it with: pip install reportlab"
            )

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm
        )

        styles = getSampleStyleSheet()

        # Custom styles
        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Heading1"],
            fontSize=24,
            spaceAfter=30,
            textColor=colors.HexColor("#00874A"),
            alignment=TA_CENTER,
        )

        subtitle_style = ParagraphStyle(
            "CustomSubtitle",
            parent=styles["Normal"],
            fontSize=12,
            textColor=colors.HexColor("#718096"),
            alignment=TA_CENTER,
            spaceAfter=20,
        )

        section_style = ParagraphStyle(
            "SectionTitle",
            parent=styles["Heading2"],
            fontSize=16,
            spaceBefore=20,
            spaceAfter=10,
            textColor=colors.HexColor("#004481"),
        )

        elements = []

        # Header
        report_title = title or f"CBAM Guard - {report_type.title()} Raporu"
        elements.append(Paragraph(report_title, title_style))
        elements.append(
            Paragraph(
                f"Oluşturulma Tarihi: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                subtitle_style,
            )
        )
        elements.append(Spacer(1, 20))

        # Content based on report type
        if report_type == "emission":
            elements.extend(self._build_emission_report(data, styles, section_style))
        elif report_type == "cbam":
            elements.extend(self._build_cbam_report(data, styles, section_style))
        elif report_type == "financial":
            elements.extend(self._build_financial_report(data, styles, section_style))
        else:
            elements.append(Paragraph(f"Rapor tipi: {report_type}", styles["Normal"]))

        # Footer
        elements.append(Spacer(1, 40))
        footer_style = ParagraphStyle(
            "Footer",
            parent=styles["Normal"],
            fontSize=10,
            textColor=colors.HexColor("#718096"),
            alignment=TA_CENTER,
        )
        elements.append(Paragraph("© 2024 CBAM Guard - Garanti BBVA", footer_style))

        doc.build(elements)
        pdf_bytes = buffer.getvalue()
        buffer.close()

        return pdf_bytes

    def _build_emission_report(self, data, styles, section_style) -> List:
        """Build emission report content"""
        from reportlab.lib import colors
        from reportlab.platypus import Paragraph, Table, TableStyle, Spacer

        elements = []

        elements.append(Paragraph("Emisyon Özeti", section_style))

        # Summary table
        summary_data = [
            ["Metrik", "Değer"],
            ["Toplam Emisyon", data.get("total_emissions", "N/A")],
            ["Scope 1", data.get("scope1", "N/A")],
            ["Scope 2", data.get("scope2", "N/A")],
            ["Scope 3", data.get("scope3", "N/A")],
        ]

        table = Table(summary_data, colWidths=[200, 200])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#00874A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 12),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F7FAFC")),
                    ("GRID", (0, 0), (-1, -1), 1, colors.HexColor("#E2E8F0")),
                ]
            )
        )

        elements.append(table)

        return elements

    def _build_cbam_report(self, data, styles, section_style) -> List:
        """Build CBAM cost report content"""
        from reportlab.lib import colors
        from reportlab.platypus import Paragraph, Table, TableStyle, Spacer

        elements = []

        elements.append(Paragraph("CBAM Maliyet Analizi", section_style))

        summary_data = [
            ["Senaryo", "Karbon Fiyatı", "Tahmini Maliyet"],
            ["Muhafazakar", "€70/tCO2", data.get("conservative_cost", "N/A")],
            ["Orta", "€90/tCO2", data.get("moderate_cost", "N/A")],
            ["Agresif", "€120/tCO2", data.get("aggressive_cost", "N/A")],
        ]

        table = Table(summary_data, colWidths=[130, 130, 130])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#004481")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 12),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F7FAFC")),
                    ("GRID", (0, 0), (-1, -1), 1, colors.HexColor("#E2E8F0")),
                ]
            )
        )

        elements.append(table)

        return elements

    def _build_financial_report(self, data, styles, section_style) -> List:
        """Build financial report content"""
        from reportlab.lib import colors
        from reportlab.platypus import Paragraph, Table, TableStyle, Spacer

        elements = []

        elements.append(Paragraph("Yeşil Finansman Özeti", section_style))

        summary_data = [
            ["Kredi Türü", "Faiz Oranı", "Vade"],
            ["Yeşil Enerji Kredisi", "%1.49", "84 ay"],
            ["Enerji Verimliliği Kredisi", "%1.79", "60 ay"],
            ["Temiz Teknoloji Kredisi", "%1.99", "72 ay"],
        ]

        table = Table(summary_data, colWidths=[150, 100, 100])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#00874A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 12),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F7FAFC")),
                    ("GRID", (0, 0), (-1, -1), 1, colors.HexColor("#E2E8F0")),
                ]
            )
        )

        elements.append(table)

        return elements

    async def generate_excel(
        self, report_type: str, data: Dict[str, Any], title: Optional[str] = None
    ) -> bytes:
        """Generate an Excel report"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ImportError(
                "openpyxl is required for Excel generation. Install it with: pip install openpyxl"
            )

        wb = Workbook()
        ws = wb.active
        ws.title = report_type.title()

        # Styles
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(
            start_color="00874A", end_color="00874A", fill_type="solid"
        )
        header_alignment = Alignment(horizontal="center", vertical="center")

        border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        # Title
        ws["A1"] = title or f"CBAM Guard - {report_type.title()} Raporu"
        ws["A1"].font = Font(bold=True, size=16, color="00874A")
        ws.merge_cells("A1:D1")

        ws["A2"] = f"Oluşturulma: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws["A2"].font = Font(italic=True, color="718096")
        ws.merge_cells("A2:D2")

        # Data based on report type
        start_row = 4

        if report_type == "emission":
            headers = ["Metrik", "Değer", "Birim", "Değişim"]
            rows = [
                ["Toplam Emisyon", data.get("total", 289000), "tCO2e", "-7.3%"],
                ["Scope 1", data.get("scope1", 145000), "tCO2e", "-5.2%"],
                ["Scope 2", data.get("scope2", 89000), "tCO2e", "-8.1%"],
                ["Scope 3", data.get("scope3", 55000), "tCO2e", "-10.5%"],
            ]
        elif report_type == "cbam":
            headers = ["Senaryo", "Karbon Fiyatı (€)", "Maliyet (€)", "Risk Seviyesi"]
            rows = [
                ["Muhafazakar", 70, data.get("conservative", 5800000), "Düşük"],
                ["Orta", 90, data.get("moderate", 7500000), "Orta"],
                ["Agresif", 120, data.get("aggressive", 9800000), "Yüksek"],
            ]
        else:
            headers = ["Kredi Türü", "Faiz (%)", "Vade (Ay)", "Uygun Yatırım"]
            rows = [
                ["Yeşil Enerji Kredisi", 1.49, 84, "Güneş, Rüzgar"],
                ["Enerji Verimliliği", 1.79, 60, "LED, Motor"],
                ["Temiz Teknoloji", 1.99, 72, "Elektrikli Araç"],
            ]

        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        # Write data
        for row_idx, row_data in enumerate(rows, start_row + 1):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border
                cell.alignment = Alignment(horizontal="center")

        # Adjust column widths
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20

        # Save to buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        excel_bytes = buffer.getvalue()
        buffer.close()

        return excel_bytes


# Global instance
export_service = ExportService()
